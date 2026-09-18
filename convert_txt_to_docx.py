"""将标准题库TXT文件转换为Word(.docx)格式"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

src_dir = r'D:\WorkXM\6a17\刷题\标准题库'
out_dir = r'D:\XIAZAI'

files = [
    'Java期末填空.txt',
    'Java期末选择判断多选.txt',
    'Python期末简答题.txt',
    'Python期末填空.txt',
    'Python期末选择判断多选.txt',
    '软件设计期末简答题.txt',
    '软件设计期末填空.txt',
    '软件设计期末选择判断多选.txt',
]

for fname in files:
    src_path = os.path.join(src_dir, fname)
    out_name = fname.replace('.txt', '.docx')
    out_path = os.path.join(out_dir, out_name)
    
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(11)
    
    with open(src_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    for line in lines:
        line = line.rstrip('\n')
        stripped = line.strip()
        
        if not stripped:
            # 空行 → 空段落（保持间距）
            doc.add_paragraph('')
            continue
        
        # 标题行 (# 开头)
        if stripped.startswith('# 题库名称'):
            p = doc.add_paragraph(stripped.lstrip('# '))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.runs[0].font
            pf.size = Pt(16)
            pf.bold = True
            pf.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            continue
        
        if stripped.startswith('## 章节'):
            p = doc.add_paragraph(stripped.lstrip('## '))
            pf = p.runs[0].font
            pf.size = Pt(14)
            pf.bold = True
            pf.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue
        
        # 题目行 (### 开头)
        if stripped.startswith('### ['):
            p = doc.add_paragraph(stripped.lstrip('### '))
            pf = p.runs[0].font
            pf.size = Pt(12)
            pf.bold = True
            continue
        
        # 选项行 (A/B/C/D 开头)
        if stripped and stripped[0] in 'ABCD' and (len(stripped) > 1 and stripped[1] in '.、 '):
            p = doc.add_paragraph(stripped)
            pf = p.runs[0].font
            pf.size = Pt(11)
            continue
        
        # 答案行
        if stripped.startswith('答案') or stripped.startswith('答案：') or stripped.startswith('答案:'):
            p = doc.add_paragraph(stripped)
            pf = p.runs[0].font
            pf.size = Pt(11)
            pf.bold = True
            pf.color.rgb = RGBColor(0x00, 0x70, 0xC0)  # 蓝色高亮答案
            continue
        
        # 解析行
        if stripped.startswith('解析') or stripped.startswith('解析：') or stripped.startswith('解析:'):
            p = doc.add_paragraph(stripped)
            pf = p.runs[0].font
            pf.size = Pt(10)
            pf.color.rgb = RGBColor(0x66, 0x66, 0x66)  # 灰色解析
            continue
        
        # 其他普通行
        p = doc.add_paragraph(stripped)
        pf = p.runs[0].font
        pf.size = Pt(11)
    
    doc.save(out_path)
    print(f'已转换: {fname} → {out_name}')

print('全部转换完成!')
